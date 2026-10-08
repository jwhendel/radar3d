#!/usr/bin/env python3
"""
Radar de Licitações 3D - Slim 3D
Varre o PNCP (Portal Nacional de Contratações Públicas) atrás de editais com
propostas em aberto que envolvam impressoras 3D, filamentos e afins.

Fluxo:
  1. Busca no PNCP por vários termos (apenas editais recebendo proposta)
  2. Remove duplicados
  3. Abre os itens de cada edital e confirma onde está o item 3D
  4. Marca o que é NOVO desde a última execução (arquivo vistos.json)
  5. Gera planilha .xlsx e, opcionalmente, envia e-mail com o resumo

Uso:
  pip install requests openpyxl
  python radar_licitacoes_3d.py

Variáveis de ambiente opcionais (para e-mail):
  SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS, EMAIL_PARA (separe vários por vírgula)
"""

import json
import os
import re
import smtplib
import sys
import time
import unicodedata
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- configuração
TERMOS_BUSCA = [
    "impressora 3d",
    "impressão 3d",
    "impressao 3d",
    "manufatura aditiva",
    "filamento pla",
    "filamento abs",
    "filamento 3d",
    "resina 3d",
    "scanner 3d",
    "laboratório maker",
    "prototipagem rápida",
]

# Regex aplicada na descrição de cada item para confirmar relevância
REGEX_3D = re.compile(
    r"(impress\w*\s*(em\s*)?3\s*d|3\s*d\s*print|manufatura\s+aditiva|"
    r"filamento|\bfdm\b|\bsla\b|resina\s+(para\s+)?(impress|3\s*d)|"
    r"scanner\s*3\s*d|prototipag)",
    re.IGNORECASE,
)

UFS_PRIORITARIAS = {"SC", "PR", "RS", "SP"}  # destaque visual, não filtra

BASE_BUSCA = "https://pncp.gov.br/api/search/"
BASE_ITENS = "https://pncp.gov.br/api/pncp/v1/orgaos/{cnpj}/compras/{ano}/{seq}/itens"
LINK_EDITAL = "https://pncp.gov.br/app/editais/{cnpj}/{ano}/{seq}"

PASTA = Path(__file__).resolve().parent
ARQ_VISTOS = PASTA / "vistos.json"
PASTA_SAIDA = PASTA / "saida"

HEADERS = {"User-Agent": "RadarLicitacoes3D/1.0 (Slim 3D)"}
SESSION = requests.Session()
SESSION.headers.update(HEADERS)


def get_json(url, params=None, tentativas=3):
    for i in range(tentativas):
        try:
            r = SESSION.get(url, params=params, timeout=40)
            if r.status_code == 204:
                return None
            r.raise_for_status()
            return r.json()
        except Exception as e:  # noqa: BLE001
            if i == tentativas - 1:
                print(f"  ! falha em {url}: {e}", file=sys.stderr)
                return None
            time.sleep(2 * (i + 1))


def sem_acento(txt):
    return unicodedata.normalize("NFKD", txt or "").encode("ascii", "ignore").decode()


# ---------------------------------------------------------------- coleta
def buscar_editais(termo, max_paginas=10):
    resultados = []
    for pagina in range(1, max_paginas + 1):
        dados = get_json(
            BASE_BUSCA,
            {
                "q": termo,
                "tipos_documento": "edital",
                "ordenacao": "-data",
                "pagina": pagina,
                "tam_pagina": 100,
                "status": "recebendo_proposta",
            },
        )
        if not dados or not dados.get("items"):
            break
        resultados.extend(dados["items"])
        if len(dados["items"]) < 100:
            break
        time.sleep(0.5)
    return resultados


def chave_edital(item):
    # item_url vem como /compras/{cnpj}/{ano}/{seq}
    partes = (item.get("item_url") or "").strip("/").split("/")
    if len(partes) >= 4:
        return partes[1], partes[2], partes[3]
    return None


def buscar_itens(cnpj, ano, seq):
    itens, pagina = [], 1
    while True:
        dados = get_json(
            BASE_ITENS.format(cnpj=cnpj, ano=ano, seq=seq),
            {"pagina": pagina, "tamanhoPagina": 500},
        )
        if not dados:
            break
        itens.extend(dados)
        if len(dados) < 500:
            break
        pagina += 1
    return itens


def classificar(edital, itens):
    objeto = f"{edital.get('title', '')} {edital.get('description', '')}"
    hits = [it for it in itens if REGEX_3D.search(sem_acento(it.get("descricao", "")))]
    no_objeto = bool(REGEX_3D.search(sem_acento(objeto)))
    if no_objeto:
        nivel = "Alta (3D no objeto)"
    elif hits:
        nivel = "Média (item 3D dentro de lote)"
    elif not itens:
        nivel = "A confirmar (itens indisponíveis)"
    else:
        nivel = None  # descartado: nenhum item 3D
    return nivel, hits


# ---------------------------------------------------------------- estado
def carregar_vistos():
    if ARQ_VISTOS.exists():
        return set(json.loads(ARQ_VISTOS.read_text(encoding="utf-8")))
    return set()


def salvar_vistos(vistos):
    ARQ_VISTOS.write_text(json.dumps(sorted(vistos), indent=1), encoding="utf-8")


# ---------------------------------------------------------------- saída
COLUNAS = [
    ("Novo?", 8), ("Relevância", 26), ("Encerra propostas", 17), ("UF", 5),
    ("Município", 18), ("Órgão", 34), ("Modalidade", 18), ("Edital", 16),
    ("Objeto", 60), ("Itens 3D encontrados", 60), ("Valor estimado itens 3D (R$)", 16),
    ("Link", 48), ("Status Slim", 16), ("Observações", 30),
]


def fmt_data(iso):
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso[:16]).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return iso


def gerar_planilha(linhas, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades"
    cab_fill = PatternFill("solid", fgColor="1F3A5F")
    for i, (nome, larg) in enumerate(COLUNAS, 1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 30

    cores = {"Alta": "D9F2D9", "Média": "FFF4CC", "A confirmar": "EDEDED"}
    for r, l in enumerate(linhas, 2):
        valores = [
            l["novo"], l["nivel"], l["encerra"], l["uf"], l["municipio"], l["orgao"],
            l["modalidade"], l["edital"], l["objeto"], l["itens_txt"], l["valor_3d"],
            l["link"], "Novo", l.get("obs", ""),
        ]
        cor = next((v for k, v in cores.items() if l["nivel"].startswith(k)), None)
        for cidx, v in enumerate(valores, 1):
            c = ws.cell(row=r, column=cidx, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if cor:
                c.fill = PatternFill("solid", fgColor=cor)
        ws.cell(row=r, column=12).hyperlink = l["link"]
        ws.cell(row=r, column=12).font = Font(color="0563C1", underline="single")
        if l["uf"] in UFS_PRIORITARIAS:
            ws.cell(row=r, column=4).font = Font(bold=True)
        ws.cell(row=r, column=11).number_format = "#,##0.00"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(len(linhas) + 1, 1)}"
    wb.save(caminho)


def enviar_email(linhas, anexo):
    host, para = os.getenv("SMTP_HOST"), os.getenv("EMAIL_PARA")
    if not host or not para:
        return
    novos = [l for l in linhas if l["novo"] == "SIM"]
    msg = EmailMessage()
    msg["Subject"] = f"Radar 3D: {len(novos)} novas licitações ({datetime.now():%d/%m})"
    msg["From"] = os.getenv("SMTP_USER")
    msg["To"] = para
    corpo = [f"{len(linhas)} oportunidades abertas, {len(novos)} novas desde ontem.\n"]
    for l in novos:
        corpo.append(
            f"[{l['nivel']}] {l['uf']} - {l['orgao']}\n"
            f"  {l['objeto'][:150]}\n  Encerra: {l['encerra']}\n  {l['link']}\n"
        )
    msg.set_content("\n".join(corpo))
    msg.add_attachment(
        anexo.read_bytes(), maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=anexo.name,
    )
    with smtplib.SMTP(host, int(os.getenv("SMTP_PORT", "587"))) as s:
        s.starttls()
        s.login(os.getenv("SMTP_USER"), os.getenv("SMTP_PASS"))
        s.send_message(msg)
    print(f"E-mail enviado para {para}")


# ---------------------------------------------------------------- principal
def main():
    print("Buscando editais no PNCP...")
    editais = {}
    for termo in TERMOS_BUSCA:
        achados = buscar_editais(termo)
        print(f"  '{termo}': {len(achados)}")
        for e in achados:
            ch = chave_edital(e)
            if ch:
                editais.setdefault(ch, e)
    print(f"{len(editais)} editais únicos. Conferindo itens...")

    vistos = carregar_vistos()
    linhas = []
    for (cnpj, ano, seq), e in editais.items():
        itens = buscar_itens(cnpj, ano, seq)
        nivel, hits = classificar(e, itens)
        if not nivel:
            continue
        controle = e.get("numero_controle_pncp") or f"{cnpj}-{ano}-{seq}"
        valor_3d = sum(float(h.get("valorTotal") or 0) for h in hits) or None
        linhas.append({
            "controle": controle,
            "novo": "NÃO" if controle in vistos else "SIM",
            "nivel": nivel,
            "encerra": fmt_data(e.get("data_fim_vigencia")),
            "encerra_iso": e.get("data_fim_vigencia") or "",
            "uf": e.get("uf", ""),
            "municipio": e.get("municipio_nome", ""),
            "orgao": e.get("orgao_nome", ""),
            "modalidade": e.get("modalidade_licitacao_nome", ""),
            "edital": e.get("title", ""),
            "objeto": e.get("description", ""),
            "itens_txt": "\n".join(
                f"Item {h.get('numeroItem')}: {h.get('descricao', '')[:120]} "
                f"(qtd {h.get('quantidade')})" for h in hits
            ),
            "valor_3d": valor_3d,
            "link": LINK_EDITAL.format(cnpj=cnpj, ano=ano, seq=seq),
        })
        time.sleep(0.3)

    def rank(n):
        return 0 if n.startswith("Alta") else 1 if n.startswith("Média") else 2
    linhas.sort(key=lambda l: (rank(l["nivel"]), l["encerra_iso"]))

    PASTA_SAIDA.mkdir(exist_ok=True)
    arq = PASTA_SAIDA / f"radar_3d_{datetime.now():%Y-%m-%d}.xlsx"
    gerar_planilha(linhas, arq)
    # JSON lido pelo painel HTML (painel_pregoes.html)
    (PASTA / "radar.json").write_text(json.dumps({
        "atualizado_em": datetime.now().isoformat(timespec="minutes"),
        "itens": linhas,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    salvar_vistos(vistos | {l["controle"] for l in linhas})
    print(f"{len(linhas)} oportunidades salvas em {arq}")
    enviar_email(linhas, arq)


if __name__ == "__main__":
    main()
